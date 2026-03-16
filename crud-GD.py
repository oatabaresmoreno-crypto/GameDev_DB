import mysql.connector
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from tkcalendar import DateEntry
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

# ============================================================
#  CONEXIÓN A LA BASE DE DATOS
# ============================================================
class DatabaseConnector:
    def __init__(self, host, user, password, database):
        self.host = host
        self.user = user
        self.password = password
        self.database = database
        self.connection = None
        self.cursor = None

    def connect(self):
        try:
            self.connection = mysql.connector.connect(
                host=self.host,
                user=self.user,
                password=self.password,
                database=self.database
            )
            self.cursor = self.connection.cursor()
            print("Conexión establecida a la base de datos")
        except mysql.connector.Error as e:
            messagebox.showerror("Error de conexión", f"Error al conectar: {e}")

    def disconnect(self):
        if self.connection:
            self.connection.close()
            print("Conexión cerrada")

    def execute_procedure(self, procedure_name, *args):
        try:
            self.cursor.callproc(procedure_name, args)
            results = []
            for result in self.cursor.stored_results():
                rows = result.fetchall()
                if rows:
                    headers = [i[0] for i in result.description]
                    results.append((headers, rows))
            self.connection.commit()
            return results
        except mysql.connector.Error as e:
            self.connection.rollback()
            messagebox.showerror("Error", f"Error en {procedure_name}: {e}")
            return None


# ============================================================
#  VENTANA PRINCIPAL
# ============================================================
class GameDevApp(tk.Tk):
    def __init__(self, db_connector):
        super().__init__()
        self.db = db_connector
        self.title("GameDev Management")
        self.geometry("1100x750")
        self.resizable(True, True)
        self._build_ui()

    def _build_ui(self):
        notebook = ttk.Notebook(self)
        notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        self.tab_vj  = VideoJuegoTab(notebook, self.db)
        self.tab_me  = MiembroEquipoTab(notebook, self.db)
        self.tab_tk  = TareaTab(notebook, self.db)
        self.tab_ag  = AssetGraficoTab(notebook, self.db)

        notebook.add(self.tab_vj,  text="🎮  Video Juegos")
        notebook.add(self.tab_me,  text="👤  Miembros")
        notebook.add(self.tab_tk,  text="📋  Tareas de Desarrollo")
        notebook.add(self.tab_ag,  text="🖼️  Assets Gráficos")


# ============================================================
#  CLASE BASE PARA CADA PESTAÑA
# ============================================================
class BaseTab(ttk.Frame):
    """Proporciona el Treeview y los métodos comunes."""

    def __init__(self, parent, db):
        super().__init__(parent)
        self.db = db
        self._build()

    def _build(self):
        # Panel izquierdo: formulario
        self.form_frame = ttk.LabelFrame(self, text="Formulario", padding=10)
        self.form_frame.pack(side=tk.LEFT, fill=tk.Y, padx=(10, 5), pady=10)

        # Panel derecho: tabla de resultados
        right = ttk.Frame(self)
        right.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(5, 10), pady=10)

        # Barra de búsqueda
        search_frame = ttk.Frame(right)
        search_frame.pack(fill=tk.X, pady=(0, 5))
        ttk.Label(search_frame, text="Buscar:").pack(side=tk.LEFT)
        self.search_var = tk.StringVar()
        ttk.Entry(search_frame, textvariable=self.search_var, width=30).pack(side=tk.LEFT, padx=5)
        ttk.Button(search_frame, text="🔍 Buscar",   command=self._search).pack(side=tk.LEFT, padx=2)
        ttk.Button(search_frame, text="🔄 Mostrar Todos", command=self._show_all).pack(side=tk.LEFT, padx=2)

        # Treeview
        tree_frame = ttk.Frame(right)
        tree_frame.pack(fill=tk.BOTH, expand=True)
        self.tree = ttk.Treeview(tree_frame, show="headings")
        vsb = ttk.Scrollbar(tree_frame, orient="vertical",   command=self.tree.yview)
        hsb = ttk.Scrollbar(tree_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        vsb.pack(side=tk.RIGHT,  fill=tk.Y)
        hsb.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree.pack(fill=tk.BOTH, expand=True)
        self.tree.bind("<<TreeviewSelect>>", self._on_select)

        # Botones CRUD
        btn_frame = ttk.Frame(right)
        btn_frame.pack(fill=tk.X, pady=5)
        ttk.Button(btn_frame, text="💾 Guardar",        command=self._insert).pack(side=tk.LEFT, padx=3)
        ttk.Button(btn_frame, text="✏️ Actualizar",     command=self._update).pack(side=tk.LEFT, padx=3)
        ttk.Button(btn_frame, text="🗑️ Eliminar",       command=self._delete).pack(side=tk.LEFT, padx=3)
        ttk.Button(btn_frame, text="🧹 Limpiar",        command=self._clear).pack(side=tk.LEFT, padx=3)
        ttk.Button(btn_frame, text="📊 Exportar Excel", command=self._export_excel).pack(side=tk.LEFT, padx=3)

        self._build_form()
        self._show_all()

    # ---- Exportar a Excel ----
    def _export_excel(self):
        # Verificar que haya datos en el Treeview
        columns = self.tree["columns"]
        rows    = [self.tree.item(i)["values"] for i in self.tree.get_children()]

        if not columns or not rows:
            messagebox.showwarning("Aviso", "No hay datos para exportar.")
            return

        # Pedir al usuario dónde guardar
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"gamedev_export_{timestamp}.xlsx"
        filepath = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            initialfile=default_name,
            title="Guardar como Excel"
        )
        if not filepath:
            return  # Usuario canceló

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Datos"

            # Estilo encabezado
            header_font  = Font(bold=True, color="FFFFFF", size=11)
            header_fill  = PatternFill("solid", fgColor="2E4057")
            header_align = Alignment(horizontal="center", vertical="center")

            # Escribir encabezados
            for col_idx, col_name in enumerate(columns, start=1):
                cell = ws.cell(row=1, column=col_idx, value=col_name)
                cell.font  = header_font
                cell.fill  = header_fill
                cell.alignment = header_align

            # Estilo filas alternadas
            fill_par  = PatternFill("solid", fgColor="EBF0F5")
            fill_impar = PatternFill("solid", fgColor="FFFFFF")

            # Escribir datos
            for row_idx, row_data in enumerate(rows, start=2):
                fill = fill_par if row_idx % 2 == 0 else fill_impar
                for col_idx, value in enumerate(row_data, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.fill = fill
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            # Ajustar ancho de columnas automáticamente
            for col in ws.columns:
                max_len = max((len(str(c.value)) for c in col if c.value), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

            # Fijar fila de encabezados
            ws.freeze_panes = "A2"

            wb.save(filepath)
            messagebox.showinfo("Éxito", f"Archivo exportado correctamente:\n{filepath}")

        except Exception as e:
            messagebox.showerror("Error al exportar", f"No se pudo guardar el archivo:\n{e}")

    # ---- Métodos que las subclases deben implementar ----
    def _build_form(self):  raise NotImplementedError
    def _get_form_values(self): raise NotImplementedError
    def _fill_form(self, values): raise NotImplementedError
    def _clear(self): raise NotImplementedError
    def _insert(self): raise NotImplementedError
    def _update(self): raise NotImplementedError
    def _delete(self): raise NotImplementedError
    def _show_all(self): raise NotImplementedError
    def _search(self): raise NotImplementedError

    # ---- Helpers ----
    def _populate_tree(self, results):
        self.tree.delete(*self.tree.get_children())
        if not results:
            return
        headers, rows = results[0]
        self.tree["columns"] = headers
        for col in headers:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=120, minwidth=60)
        for row in rows:
            self.tree.insert("", tk.END, values=row)

    def _on_select(self, event):
        sel = self.tree.selection()
        if sel:
            self._fill_form(self.tree.item(sel[0])["values"])

    def _lbl_entry(self, label, row, width=25, state="normal"):
        ttk.Label(self.form_frame, text=label).grid(row=row, column=0, sticky=tk.W, pady=3, padx=(0,8))
        e = ttk.Entry(self.form_frame, width=width, state=state)
        e.grid(row=row, column=1, sticky=tk.W, pady=3)
        return e

    def _lbl_combo(self, label, row, values, width=22):
        ttk.Label(self.form_frame, text=label).grid(row=row, column=0, sticky=tk.W, pady=3, padx=(0,8))
        c = ttk.Combobox(self.form_frame, values=values, width=width, state="readonly")
        c.grid(row=row, column=1, sticky=tk.W, pady=3)
        return c

    def _lbl_date(self, label, row):
        ttk.Label(self.form_frame, text=label).grid(row=row, column=0, sticky=tk.W, pady=3, padx=(0,8))
        d = DateEntry(self.form_frame, width=22, background="darkblue",
                      foreground="white", borderwidth=2, date_pattern="yyyy-mm-dd")
        d.grid(row=row, column=1, sticky=tk.W, pady=3)
        return d

    def _set_entry(self, entry, value):
        if isinstance(entry, DateEntry):
            try:
                entry.set_date(str(value)) if value else None
            except Exception:
                pass
        elif isinstance(entry, ttk.Combobox):
            entry.set(str(value) if value else "")
        else:
            entry.config(state="normal")
            entry.delete(0, tk.END)
            entry.insert(0, str(value) if value is not None else "")

    def _get(self, entry):
        if isinstance(entry, DateEntry):
            return entry.get_date().strftime("%Y-%m-%d")
        v = entry.get().strip()
        return v if v else None


# ============================================================
#  TAB 1 — VIDEO JUEGOS
# ============================================================
class VideoJuegoTab(BaseTab):

    def _build_form(self):
        f = self.form_frame
        self.e_codigo      = self._lbl_entry("Código Único:",            0)
        self.e_titulo      = self._lbl_entry("Título:",                   1)
        self.e_genero      = self._lbl_combo("Género:",                   2, ["accion","aventura","estrategia","rol"])
        self.e_plataformas = self._lbl_combo("Plataformas:",              3, ["PC","Consolas","Moviles"])
        self.e_clas_tipo   = self._lbl_combo("Clasificación (tipo):",     4, ["ESRB","PEGI"])
        self.e_clas_valor  = self._lbl_entry("Clasificación (valor):",    5)
        self.e_f_inicio    = self._lbl_date ("Fecha Inicio:",             6)
        self.e_f_lanz      = self._lbl_date ("Fecha Lanzamiento:",        7)
        self.e_presupuesto = self._lbl_entry("Presupuesto:",              8)
        self.e_motor       = self._lbl_entry("Motor Gráfico:",            9)
        self.e_equipo      = self._lbl_entry("Equipo de Desarrollo:",    10)
        self.e_estado      = self._lbl_combo("Estado:",                  11,
            ["pre-produccion","en desarrollo","beta","lanzado","cancelado"])

    def _get_form_values(self):
        return (
            self._get(self.e_codigo),
            self._get(self.e_titulo),
            self._get(self.e_genero),
            self._get(self.e_plataformas),
            self._get(self.e_clas_tipo),
            self._get(self.e_clas_valor),
            self._get(self.e_f_inicio),
            self._get(self.e_f_lanz),
            self._get(self.e_presupuesto),
            self._get(self.e_motor),
            self._get(self.e_equipo),
            self._get(self.e_estado),
        )

    def _fill_form(self, v):
        fields = [self.e_codigo, self.e_titulo, self.e_genero, self.e_plataformas,
                  self.e_clas_tipo, self.e_clas_valor, self.e_f_inicio, self.e_f_lanz,
                  self.e_presupuesto, self.e_motor, self.e_equipo, self.e_estado]
        for widget, val in zip(fields, v):
            self._set_entry(widget, val)

    def _clear(self):
        for w in [self.e_codigo, self.e_titulo, self.e_genero, self.e_plataformas,
                  self.e_clas_tipo, self.e_clas_valor, self.e_presupuesto,
                  self.e_motor, self.e_equipo, self.e_estado]:
            self._set_entry(w, "")

    def _insert(self):
        v = self._get_form_values()
        if not v[0] or not v[1]:
            messagebox.showwarning("Aviso", "Código y Título son obligatorios.")
            return
        r = self.db.execute_procedure("sp_InsertVideojuego", *v)
        if r is not None:
            messagebox.showinfo("Éxito", "Videojuego guardado correctamente.")
            self._clear(); self._show_all()

    def _update(self):
        v = self._get_form_values()
        if not v[0]:
            messagebox.showwarning("Aviso", "Selecciona un videojuego primero.")
            return
        r = self.db.execute_procedure("sp_UpdateVideojuego", *v)
        if r is not None:
            messagebox.showinfo("Éxito", "Videojuego actualizado correctamente.")
            self._show_all()

    def _delete(self):
        codigo = self._get(self.e_codigo)
        if not codigo:
            messagebox.showwarning("Aviso", "Selecciona un videojuego primero.")
            return
        if messagebox.askyesno("Confirmar", f"¿Eliminar el videojuego '{codigo}'?"):
            r = self.db.execute_procedure("sp_DeleteVideojuego", codigo)
            if r is not None:
                messagebox.showinfo("Éxito", "Videojuego eliminado.")
                self._clear(); self._show_all()

    def _show_all(self):
        r = self.db.execute_procedure("sp_GetAllVideojuegos")
        self._populate_tree(r)

    def _search(self):
        term = self.search_var.get().strip()
        if not term:
            self._show_all(); return
        r = self.db.execute_procedure("sp_SearchVideojuegos", term)
        self._populate_tree(r)


# ============================================================
#  TAB 2 — MIEMBROS DEL EQUIPO
# ============================================================
class MiembroEquipoTab(BaseTab):

    def _build_form(self):
        self.e_codigo      = self._lbl_entry("Código Empleado:",       0)
        self.e_nombres     = self._lbl_entry("Nombres:",               1)
        self.e_apellidos   = self._lbl_entry("Apellidos:",             2)
        self.e_especialidad= self._lbl_combo("Especialidad:",          3,
            ["programacion","arte","diseno","musica","testing"])
        self.e_habilidades = self._lbl_entry("Habilidades:",           4, width=30)
        self.e_experiencia = self._lbl_entry("Experiencia Previa:",    5, width=30)

    def _get_form_values(self):
        return (
            self._get(self.e_codigo),
            self._get(self.e_nombres),
            self._get(self.e_apellidos),
            self._get(self.e_especialidad),
            self._get(self.e_habilidades),
            self._get(self.e_experiencia),
        )

    def _fill_form(self, v):
        widgets = [self.e_codigo, self.e_nombres, self.e_apellidos,
                   self.e_especialidad, self.e_habilidades, self.e_experiencia]
        # El treeview puede traer más columnas (total_proyectos, creado_en), sólo usamos las primeras 6
        for w, val in zip(widgets, v):
            self._set_entry(w, val)

    def _clear(self):
        for w in [self.e_codigo, self.e_nombres, self.e_apellidos,
                  self.e_especialidad, self.e_habilidades, self.e_experiencia]:
            self._set_entry(w, "")

    def _insert(self):
        v = self._get_form_values()
        if not v[0] or not v[1]:
            messagebox.showwarning("Aviso", "Código y Nombres son obligatorios.")
            return
        r = self.db.execute_procedure("sp_InsertMiembroEquipo", *v)
        if r is not None:
            messagebox.showinfo("Éxito", "Miembro guardado correctamente.")
            self._clear(); self._show_all()

    def _update(self):
        v = self._get_form_values()
        if not v[0]:
            messagebox.showwarning("Aviso", "Selecciona un miembro primero.")
            return
        r = self.db.execute_procedure("sp_UpdateMiembroEquipo", *v)
        if r is not None:
            messagebox.showinfo("Éxito", "Miembro actualizado correctamente.")
            self._show_all()

    def _delete(self):
        codigo = self._get(self.e_codigo)
        if not codigo:
            messagebox.showwarning("Aviso", "Selecciona un miembro primero.")
            return
        if messagebox.askyesno("Confirmar", f"¿Eliminar al miembro '{codigo}'?"):
            r = self.db.execute_procedure("sp_DeleteMiembroEquipo", codigo)
            if r is not None:
                messagebox.showinfo("Éxito", "Miembro eliminado.")
                self._clear(); self._show_all()

    def _show_all(self):
        r = self.db.execute_procedure("sp_GetAllMiembrosEquipo")
        self._populate_tree(r)

    def _search(self):
        term = self.search_var.get().strip()
        if not term:
            self._show_all(); return
        r = self.db.execute_procedure("sp_SearchMiembrosEquipo", term)
        self._populate_tree(r)


# ============================================================
#  TAB 3 — TAREAS DE DESARROLLO
# ============================================================
class TareaTab(BaseTab):

    def _build_form(self):
        self.e_codigo      = self._lbl_entry("Código Tarea:",          0)
        self.e_videojuego  = self._lbl_entry("Código Videojuego:",     1)
        self.e_modulo      = self._lbl_combo("Módulo:",                2,
            ["jugabilidad","graficos","sonido","ia"])
        self.e_descripcion = self._lbl_entry("Descripción:",           3, width=30)
        self.e_prioridad   = self._lbl_combo("Prioridad:",             4,
            ["baja","media","alta","critica"])
        self.e_estado      = self._lbl_combo("Estado:",                5,
            ["pendiente","en proceso","revision","completada"])
        self.e_responsable = self._lbl_entry("Responsable (código):",  6)
        self.e_f_asignacion= self._lbl_date ("Fecha Asignación:",      7)
        self.e_f_limite    = self._lbl_date ("Fecha Límite:",          8)
        self.e_avance      = self._lbl_entry("% Avance:",              9)
        self.e_dependencias= self._lbl_entry("Dependencias:",         10, width=30)

    def _get_form_values(self):
        return (
            self._get(self.e_codigo),
            self._get(self.e_videojuego),
            self._get(self.e_modulo),
            self._get(self.e_descripcion),
            self._get(self.e_prioridad),
            self._get(self.e_estado),
            self._get(self.e_responsable),
            self._get(self.e_f_asignacion),
            self._get(self.e_f_limite),
            self._get(self.e_avance) or "0",
            self._get(self.e_dependencias),
        )

    def _fill_form(self, v):
        widgets = [self.e_codigo, self.e_videojuego, self.e_modulo,
                   self.e_descripcion, self.e_prioridad, self.e_estado,
                   self.e_responsable, self.e_f_asignacion, self.e_f_limite,
                   self.e_avance, self.e_dependencias]
        for w, val in zip(widgets, v):
            self._set_entry(w, val)

    def _clear(self):
        for w in [self.e_codigo, self.e_videojuego, self.e_modulo,
                  self.e_descripcion, self.e_prioridad, self.e_estado,
                  self.e_responsable, self.e_avance, self.e_dependencias]:
            self._set_entry(w, "")

    def _insert(self):
        v = self._get_form_values()
        if not v[0] or not v[1]:
            messagebox.showwarning("Aviso", "Código de tarea y Videojuego son obligatorios.")
            return
        r = self.db.execute_procedure("sp_InsertTarea", *v)
        if r is not None:
            messagebox.showinfo("Éxito", "Tarea guardada correctamente.")
            self._clear(); self._show_all()

    def _update(self):
        vals = self._get_form_values()
        if not vals[0]:
            messagebox.showwarning("Aviso", "Selecciona una tarea primero.")
            return
        # sp_UpdateTarea no recibe codigo_videojuego, sólo los campos modificables
        update_args = (vals[0],) + vals[2:]   # codigo, modulo, desc, prioridad, estado, resp, f_asig, f_lim, avance, dep
        r = self.db.execute_procedure("sp_UpdateTarea", *update_args)
        if r is not None:
            messagebox.showinfo("Éxito", "Tarea actualizada correctamente.")
            self._show_all()

    def _delete(self):
        codigo = self._get(self.e_codigo)
        if not codigo:
            messagebox.showwarning("Aviso", "Selecciona una tarea primero.")
            return
        if messagebox.askyesno("Confirmar", f"¿Eliminar la tarea '{codigo}'?"):
            r = self.db.execute_procedure("sp_DeleteTarea", codigo)
            if r is not None:
                messagebox.showinfo("Éxito", "Tarea eliminada.")
                self._clear(); self._show_all()

    def _show_all(self):
        # Mostramos todas las tareas usando búsqueda vacía (muestra todo)
        r = self.db.execute_procedure("sp_SearchTareas", "")
        self._populate_tree(r)

    def _search(self):
        term = self.search_var.get().strip()
        r = self.db.execute_procedure("sp_SearchTareas", term)
        self._populate_tree(r)


# ============================================================
#  TAB 4 — ASSETS GRÁFICOS
# ============================================================
class AssetGraficoTab(BaseTab):

    def _build_form(self):
        self.e_codigo        = self._lbl_entry("Código Asset:",             0)
        self.e_nombre        = self._lbl_entry("Nombre:",                   1, width=30)
        self.e_tipo          = self._lbl_combo("Tipo:",                     2,
            ["personaje","escenario","objeto","interfaz"])
        self.e_descripcion   = self._lbl_entry("Descripción:",              3, width=30)
        self.e_formato       = self._lbl_entry("Formato (PNG/FBX/…):",      4)
        self.e_resolucion    = self._lbl_entry("Resolución:",               5)
        self.e_artista       = self._lbl_entry("Artista (código):",         6)
        self.e_version       = self._lbl_entry("Versión:",                  7)
        self.e_f_modif       = self._lbl_date ("Fecha Modificación:",       8)
        self.e_requisitos    = self._lbl_entry("Requisitos Técnicos:",      9, width=30)
        self.e_estado_apr    = self._lbl_combo("Estado Aprobación:",       10,
            ["pendiente","en revision","aprobado","rechazado"])
        self.e_videojuego    = self._lbl_entry("Código Videojuego:",       11)

    def _get_form_values(self):
        return (
            self._get(self.e_codigo),
            self._get(self.e_nombre),
            self._get(self.e_tipo),
            self._get(self.e_descripcion),
            self._get(self.e_formato),
            self._get(self.e_resolucion),
            self._get(self.e_artista),
            self._get(self.e_version) or "1.0",
            self._get(self.e_f_modif),
            self._get(self.e_requisitos),
            self._get(self.e_estado_apr),
            self._get(self.e_videojuego),
        )

    def _fill_form(self, v):
        widgets = [self.e_codigo, self.e_nombre, self.e_tipo,
                   self.e_descripcion, self.e_formato, self.e_resolucion,
                   self.e_artista, self.e_version, self.e_f_modif,
                   self.e_requisitos, self.e_estado_apr, self.e_videojuego]
        for w, val in zip(widgets, v):
            self._set_entry(w, val)

    def _clear(self):
        for w in [self.e_codigo, self.e_nombre, self.e_tipo,
                  self.e_descripcion, self.e_formato, self.e_resolucion,
                  self.e_artista, self.e_version, self.e_requisitos,
                  self.e_estado_apr, self.e_videojuego]:
            self._set_entry(w, "")

    def _insert(self):
        v = self._get_form_values()
        if not v[0] or not v[1]:
            messagebox.showwarning("Aviso", "Código y Nombre son obligatorios.")
            return
        r = self.db.execute_procedure("sp_InsertAssetGrafico", *v)
        if r is not None:
            messagebox.showinfo("Éxito", "Asset guardado correctamente.")
            self._clear(); self._show_all()

    def _update(self):
        v = self._get_form_values()
        if not v[0]:
            messagebox.showwarning("Aviso", "Selecciona un asset primero.")
            return
        r = self.db.execute_procedure("sp_UpdateAssetGrafico", *v)
        if r is not None:
            messagebox.showinfo("Éxito", "Asset actualizado correctamente.")
            self._show_all()

    def _delete(self):
        codigo = self._get(self.e_codigo)
        if not codigo:
            messagebox.showwarning("Aviso", "Selecciona un asset primero.")
            return
        if messagebox.askyesno("Confirmar", f"¿Eliminar el asset '{codigo}'?"):
            r = self.db.execute_procedure("sp_DeleteAssetGrafico", codigo)
            if r is not None:
                messagebox.showinfo("Éxito", "Asset eliminado.")
                self._clear(); self._show_all()

    def _show_all(self):
        r = self.db.execute_procedure("sp_GetAllAssetsGraficos")
        self._populate_tree(r)

    def _search(self):
        term = self.search_var.get().strip()
        if not term:
            self._show_all(); return
        r = self.db.execute_procedure("sp_SearchAssetsGraficos", term)
        self._populate_tree(r)


# ============================================================
#  PUNTO DE ENTRADA
# ============================================================
if __name__ == "__main__":
    db = DatabaseConnector(
        host="localhost",
        user="root",
        password="",          # ← pon tu contraseña aquí
        database="gamedev"
    )
    db.connect()

    app = GameDevApp(db)
    app.mainloop()

    db.disconnect()